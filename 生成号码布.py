#!/usr/bin/env python3
"""从 Excel 名单按学校、组别和队伍批量制作双面号码布。"""

from __future__ import annotations

import argparse
import csv
import io
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ByteStringObject, ContentStream
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


PROJECT = Path(__file__).resolve().parent
INPUT_XLSX: Path | None = None
OUTPUT_ROOT = PROJECT / "output"
TITLE_FONT: Path | None = PROJECT / "fonts/优设标题黑_猫啃网.ttf"
FALLBACK_FONT: Path | None = None
KUN_GLYPH: Path | None = None

# 女子组、混合组颜色来自本次赛事的自制模板；示例模板仅用于演示。
# 坐标单位为 PDF pt，原点在左下角。
TEMPLATES = {
    "男子组": {
        "path": PROJECT / "private_assets/校园赛号码布_男子组_官方.pdf",
        "placeholders": ("孙悟空", "兰州大学"),
        "front_page": 1,
        "center_x": 408.5,
        "name_center_y": 339.4,
        "name_size": 164.15,
        "name_max_width": 650.0,
        "name_min_size": 70.0,
        "two_line_size": 90.0,
        "two_line_centers_y": (382.0, 303.0),
        "school_center_y": 224.8,
        "school_size": 59.73,
        "school_max_width": 420.0,
        "school_min_size": 32.0,
        "school_color": (0.58, 0.269, 0.137),
        "clear_text_indexes": (0, 1),
        "text_object_count": 3,
        "static_text_checks": ("承办单位：",),
    },
}
TEMPLATES["女子组"] = {
    **TEMPLATES["男子组"],
    "path": PROJECT / "private_assets/校园赛号码布_女子组_自制深蓝.pdf",
    "school_color": (51 / 255, 97 / 255, 126 / 255),
}
TEMPLATES["混合组"] = {
    **TEMPLATES["男子组"],
    "path": PROJECT / "private_assets/校园赛号码布_混合组_自制青色.pdf",
    "school_color": (33 / 255, 136 / 255, 153 / 255),
}

HEADERS = ("序号", "组别", "组队名称", "学生姓名", "性别", "学生所属学院名称")
MANIFEST_HEADERS = ("Excel行号", "原序号", "学校", "组别", "队伍", "姓名", "性别", "状态", "PDF路径", "备注")
FONT_NAMES = {}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value).strip()) if value is not None else ""


def display_label(value: str) -> str:
    return (value.replace("香港科技大学（广州）", "香港科技大学(广州)")
                 .replace("香港科技大学（HKUST)", "香港科技大学(HKUST)")
                 .replace("香港科技大学（HKUST）", "香港科技大学(HKUST)"))


def safe_component(value: str) -> str:
    safe = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", value).strip(" .")
    return safe or "未命名"


def read_roster() -> list[dict[str, str | int]]:
    if INPUT_XLSX is None:
        return [
            {"序号": str(index), "组别": group, "组队名称": "演示队",
             "学生姓名": name, "性别": sex, "学生所属学院名称": "示例大学", "Excel行号": index + 1}
            for index, (group, name, sex) in enumerate((
                ("男子组", "示例甲", "男"), ("女子组", "示例乙", "女"),
                ("混合组", "示例丙", "男")), start=1)
        ]
    if not INPUT_XLSX.is_file():
        raise FileNotFoundError(f"找不到名单：{INPUT_XLSX}")
    book = load_workbook(INPUT_XLSX, read_only=True, data_only=True)
    try:
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows)]
        missing = set(HEADERS) - set(headers)
        if missing:
            raise ValueError(f"名单缺少列：{', '.join(sorted(missing))}")
        positions = {name: headers.index(name) for name in HEADERS}
        result = []
        for excel_row, values in enumerate(rows, start=2):
            if not any(value is not None and str(value).strip() for value in values):
                continue
            record = {name: clean(values[positions[name]]) if positions[name] < len(values) else "" for name in HEADERS}
            record["学生所属学院名称"] = display_label(record["学生所属学院名称"])
            record["组队名称"] = display_label(record["组队名称"])
            record["Excel行号"] = excel_row
            result.append(record)
        return result
    finally:
        book.close()


def register_fonts() -> None:
    FONT_NAMES["title"] = ("Helvetica-Bold", set(range(32, 127)))
    if TITLE_FONT is not None:
        if not TITLE_FONT.is_file():
            raise FileNotFoundError(f"缺少标题字体：{TITLE_FONT}")
        font = TTFont("title", str(TITLE_FONT))
        pdfmetrics.registerFont(font)
        FONT_NAMES["title"] = ("title", font.face.charToGlyph)
    if FALLBACK_FONT is not None:
        if not FALLBACK_FONT.is_file():
            raise FileNotFoundError(f"缺少备用字体：{FALLBACK_FONT}")
        font = TTFont("fallback", str(FALLBACK_FONT))
        pdfmetrics.registerFont(font)
        FONT_NAMES["fallback"] = ("fallback", font.face.charToGlyph)


def font_for(character: str) -> str:
    code = ord(character)
    for label in ("title", "fallback"):
        if label in FONT_NAMES:
            name, coverage = FONT_NAMES[label]
            if code in coverage:
                return name
    raise ValueError(f"字体无法显示字符 {character!r}；请配置包含该字的备用字体")


def character_width(character: str, size: float, custom_glyph: bool = False) -> float:
    if custom_glyph and character == "堃" and KUN_GLYPH is not None:
        return pdfmetrics.stringWidth("梁", FONT_NAMES["title"][0], size)
    return pdfmetrics.stringWidth(character, font_for(character), size)


def measured_width(value: str, size: float, custom_glyph: bool = False) -> float:
    return sum(character_width(character, size, custom_glyph) for character in value) + max(0, len(value) - 1) * size * 0.1


def fitted_size(lines: list[str], original: float, max_width: float, minimum: float, custom_glyph: bool = False) -> float:
    widest = max(measured_width(line, original, custom_glyph) for line in lines)
    size = min(original, original * max_width / widest) if widest else original
    if size < minimum:
        raise ValueError(f"文字过长，适配字号仅 {size:.1f} pt；请人工排版：{' / '.join(lines)}")
    return size


def draw_centered(c: canvas.Canvas, value: str, size: float, center_x: float, center_y: float,
                  color: tuple[float, float, float], custom_glyph: bool = False) -> None:
    width = measured_width(value, size, custom_glyph)
    # 0.233 是随附字体在原 Illustrator 样张中的可见字框中心相对基线比例。
    baseline = center_y - size * 0.233
    if custom_glyph and "堃" in value and KUN_GLYPH is not None:
        if not KUN_GLYPH.is_file():
            raise FileNotFoundError(f"缺少定稿字形：{KUN_GLYPH}")
        x = center_x - width / 2
        c.setFillColorRGB(*color)
        for character in value:
            if character == "堃":
                c.saveState()
                hidden = c.beginText(x, baseline)
                hidden.setFont(font_for(character), size)
                hidden.setTextRenderMode(3)
                hidden.textOut(character)
                c.drawText(hidden)
                c.restoreState()
                c.drawImage(str(KUN_GLYPH), x, baseline - size * 0.009,
                            width=character_width(character, size, True), height=size * 0.72, mask="auto")
            else:
                c.setFont(font_for(character), size)
                c.drawString(x, baseline, character)
            x += character_width(character, size, True) + size * 0.1
        return
    text = c.beginText(center_x - width / 2, baseline)
    text.setFillColorRGB(*color)
    text.setCharSpace(size * 0.1)
    for character in value:
        text.setFont(font_for(character), size)
        text.textOut(character)
    c.drawText(text)


def overlay_for(name: str, school: str, page_width: float, page_height: float, spec: dict) -> bytes:
    stream = io.BytesIO()
    c = canvas.Canvas(stream, pagesize=(page_width, page_height), pageCompression=1)
    if " " in name and measured_width(name, spec["name_size"], True) > spec["name_max_width"]:
        first, second = name.split(" ", 1)
        lines = [first, second]
        size = fitted_size(lines, spec["two_line_size"], spec["name_max_width"], spec["name_min_size"], True)
        for line, center_y in zip(lines, spec["two_line_centers_y"]):
            draw_centered(c, line, size, spec["center_x"], center_y, (1, 1, 1), True)
    else:
        size = fitted_size([name], spec["name_size"], spec["name_max_width"], spec["name_min_size"], True)
        draw_centered(c, name, size, spec["center_x"], spec["name_center_y"], (1, 1, 1), True)
    school_size = fitted_size([school], spec["school_size"], spec["school_max_width"], spec["school_min_size"])
    draw_centered(c, school, school_size, spec["center_x"], spec["school_center_y"], spec["school_color"])
    c.save()
    return stream.getvalue()


def make_demo_template(path: Path, group: str, color: tuple[float, float, float]) -> None:
    """画一份无赛事标志、无真实姓名的双页示例模板。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=(844.779, 595.761), pageCompression=1)
    c.setFont("Helvetica", 17)
    c.drawString(55, 535, "DEMO BIB - BACK")
    for label, y in (("Runner name:", 445), ("Emergency contact:", 365), ("Phone:", 285)):
        c.setStrokeColorRGB(.65, .65, .65)
        c.roundRect(55, y - 20, 735, 55, 9, fill=0, stroke=1)
        c.drawString(72, y, label)
    c.drawString(55, 170, "Example only. Use licensed assets for real events.")
    c.showPage()

    c.setFillColorRGB(1, 1, 1)
    c.rect(0, 0, 844.779, 595.761, fill=1, stroke=0)
    c.setFillColorRGB(*color)
    c.rect(25, 70, 794, 405, fill=1, stroke=0)
    c.setFillColorRGB(1, 1, 1)
    c.roundRect(190, 190, 437, 74, 37, fill=1, stroke=0)
    for x in (48, 797):
        c.setFillColorRGB(.05, .05, .05)
        c.circle(x, 35, 10, fill=1, stroke=0)
    # 以下前两个文字对象会由脚本清除并替换为名单中的姓名和学校。
    c.setFont("Helvetica", 16)
    c.setFillColorRGB(1, 1, 1)
    c.drawString(340, 340, "PLACEHOLDER NAME")
    c.setFillColorRGB(*color)
    c.drawString(340, 220, "PLACEHOLDER SCHOOL")
    c.setFont("Helvetica-Bold", 24)
    c.setFillColorRGB(.12, .12, .12)
    c.drawString(310, 515, "DEMO TEMPLATE")
    c.setFont("Helvetica-Bold", 22)
    c.setFillColorRGB(1, 1, 1)
    c.drawString(370, 100, {"男子组": "MEN", "女子组": "WOMEN", "混合组": "MIXED"}[group])
    c.save()


def prepared_template(spec: dict) -> bytes:
    path = spec["path"]
    if not path.is_file():
        raise FileNotFoundError(f"缺少模板：{path}")
    reader = PdfReader(path)
    if len(reader.pages) != 2 or spec["front_page"] not in (0, 1):
        raise ValueError(f"模板页数或正面页号不符：{path}")
    front = reader.pages[spec["front_page"]]
    original_text = front.extract_text()
    if not all(placeholder in original_text for placeholder in spec["placeholders"]):
        raise ValueError(f"模板示例文字与配置不符：{path}")
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    front = writer.pages[spec["front_page"]]
    contents = ContentStream(front.get_contents(), writer)
    text_objects = [args for args, operator in contents.operations if operator == b"Tj"]
    if len(text_objects) != spec["text_object_count"]:
        raise ValueError(f"模板文字对象数量已变化，请重新核对：{path}")
    for index in spec["clear_text_indexes"]:
        text_objects[index][0] = ByteStringObject(b"")
    front.replace_contents(contents)
    if any(placeholder in front.extract_text() for placeholder in spec["placeholders"]):
        raise ValueError(f"未能清除模板示例姓名或学校：{path}")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_pdf(blank_template: bytes, spec: dict, name: str, school: str, output: Path) -> None:
    reader = PdfReader(io.BytesIO(blank_template))
    writer = PdfWriter()
    writer.append_pages_from_reader(reader)
    page = writer.pages[spec["front_page"]]
    overlay = PdfReader(io.BytesIO(overlay_for(name, school, float(page.mediabox.width), float(page.mediabox.height), spec)))
    page.merge_page(overlay.pages[0])
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(".pdf.tmp")
    with temporary.open("wb") as handle:
        writer.write(handle)
    temporary.replace(output)
    check = PdfReader(output)
    text = check.pages[spec["front_page"]].extract_text()
    compact_text = re.sub(r"\s+", "", text)
    if (len(check.pages) != 2 or re.sub(r"\s+", "", name) not in compact_text
            or re.sub(r"\s+", "", school) not in compact_text
            or any(required not in text for required in spec["static_text_checks"])):
        output.unlink()
        raise ValueError("生成后 PDF 页数或姓名、学校、静态文字核对失败")


def fresh_batch_dir() -> Path:
    base = datetime.now().strftime("%Y%m%d_%H%M%S")
    for suffix in range(1000):
        path = OUTPUT_ROOT / (base if suffix == 0 else f"{base}_{suffix:03d}")
        try:
            path.mkdir(parents=True)
            return path
        except FileExistsError:
            continue
    raise RuntimeError("无法创建新的生成批次目录")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="按学校、组别、队伍生成双面号码布 PDF")
    parser.add_argument("--input", type=Path, help="Excel 名单路径；不填则运行匿名示例")
    parser.add_argument("--output", type=Path, default=PROJECT / "output", help="输出根目录")
    parser.add_argument("--templates-dir", type=Path, help="已获授权的本次赛事三份 PDF 模板所在目录")
    parser.add_argument("--title-font", type=Path, help="标题字体 TTF 路径")
    parser.add_argument("--fallback-font", type=Path, help="备用字体 TTF 路径")
    parser.add_argument("--kun-glyph", type=Path, help="已定稿的“堃”字形 PNG 路径")
    args = parser.parse_args(argv)
    if args.templates_dir and not args.input:
        parser.error("使用赛事模板时必须提供 --input")

    global INPUT_XLSX, OUTPUT_ROOT, TITLE_FONT, FALLBACK_FONT, KUN_GLYPH
    INPUT_XLSX = args.input.expanduser().resolve() if args.input else None
    OUTPUT_ROOT = args.output.expanduser().resolve()
    TITLE_FONT = args.title_font.expanduser().resolve() if args.title_font else PROJECT / "fonts/优设标题黑_猫啃网.ttf"
    FALLBACK_FONT = args.fallback_font.expanduser().resolve() if args.fallback_font else None
    KUN_GLYPH = args.kun_glyph.expanduser().resolve() if args.kun_glyph else None

    roster = read_roster()
    if args.templates_dir and any("堃" in person["学生姓名"] for person in roster):
        if KUN_GLYPH is None or FALLBACK_FONT is None:
            parser.error("名单含“堃”；使用本次赛事模板时请提供 --kun-glyph 和 --fallback-font")
    register_fonts()
    if args.templates_dir:
        template_dir = args.templates_dir.expanduser().resolve()
        for spec in TEMPLATES.values():
            spec["path"] = template_dir / spec["path"].name
    else:
        demo_dir = OUTPUT_ROOT / "_demo_templates"
        for group, spec in TEMPLATES.items():
            spec["path"] = demo_dir / f"{group}.pdf"
            spec["placeholders"] = ("PLACEHOLDER NAME", "PLACEHOLDER SCHOOL")
            spec["text_object_count"] = 4
            spec["static_text_checks"] = ("DEMO TEMPLATE", {"男子组": "MEN", "女子组": "WOMEN", "混合组": "MIXED"}[group])
            make_demo_template(spec["path"], group, spec["school_color"])
    prepared = {}
    template_errors = {}
    for group, spec in TEMPLATES.items():
        if spec is None:
            continue
        try:
            prepared[group] = prepared_template(spec)
        except Exception as exc:
            template_errors[group] = str(exc)
    batch = fresh_batch_dir()
    records = []
    seen = set()
    team_sizes = defaultdict(int)
    for person in roster:
        row = person["Excel行号"]
        group = person["组别"]
        team = person["组队名称"]
        name = person["学生姓名"]
        school = person["学生所属学院名称"]
        sex = person["性别"]
        entry = {"Excel行号": row, "原序号": person["序号"], "学校": school, "组别": group,
                 "队伍": team, "姓名": name, "性别": sex, "状态": "", "PDF路径": "", "备注": ""}
        missing = [label for label, value in (("学校", school), ("组别", group), ("队伍", team), ("姓名", name), ("性别", sex)) if not value]
        key = (school, group, team, name.casefold())
        if missing:
            entry["状态"] = "缺少必填字段"
            entry["备注"] = "、".join(missing)
        elif key in seen:
            entry["状态"] = "重复记录"
            entry["备注"] = "同校、同组、同队已有相同姓名"
        elif (group == "男子组" and sex != "男") or (group == "女子组" and sex != "女"):
            entry["状态"] = "组别性别冲突"
            entry["备注"] = f"{group}的性别记录为{sex}，请核对名单"
        elif group not in TEMPLATES:
            entry["状态"] = "组别未知"
        else:
            seen.add(key)
            team_sizes[(school, group, team)] += 1
            spec = TEMPLATES[group]
            if spec is None:
                entry["状态"] = "待模板"
                entry["备注"] = "该组别尚无正式模板"
            elif group in template_errors:
                entry["状态"] = "模板错误"
                entry["备注"] = template_errors[group]
            else:
                folder = batch / safe_component(school) / safe_component(group) / safe_component(team)
                output = folder / f"{safe_component(name)}.pdf"
                if output.exists():
                    output = folder / f"{safe_component(name)}_Excel行{row}.pdf"
                try:
                    make_pdf(prepared[group], spec, name, school, output)
                    entry["状态"] = "已生成"
                    entry["PDF路径"] = str(output.relative_to(batch))
                    notes = []
                    if "堃" in name and KUN_GLYPH is not None:
                        notes.append("堃使用定稿字形")
                    if "fallback" in FONT_NAMES and any(
                        font_for(character) == FONT_NAMES["fallback"][0]
                        for character in name.replace("堃", "") + school
                    ):
                        notes.append("个别字符使用备用字体")
                    entry["备注"] = "；".join(notes)
                except Exception as exc:
                    entry["状态"] = "生成失败"
                    entry["备注"] = str(exc)
        records.append(entry)

    with (batch / "生成清单.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=MANIFEST_HEADERS)
        writer.writeheader()
        writer.writerows(records)
    counts = Counter(record["状态"] for record in records)
    warnings = [f"{school} / {group} / {team}：{count} 人（不是 5 人）"
                for (school, group, team), count in team_sizes.items() if count != 5]
    group_counts = Counter(person["组别"] for person in roster)
    report = [f"数据源：{INPUT_XLSX or '内置匿名示例'}", f"报名行数：{len(roster)}",
              "组别人数：" + "，".join(f"{group} {count}" for group, count in group_counts.items()), "处理结果："]
    report += [f"- {status}：{count} 人" for status, count in counts.items()]
    report += ["", "队伍人数提醒："] + ([f"- {item}" for item in warnings] if warnings else ["- 无"])
    (batch / "运行结果.txt").write_text("\n".join(report) + "\n", encoding="utf-8")
    print("\n".join(report))
    print(f"输出位置：{batch}")
    return 0 if counts.get("已生成", 0) == len(roster) else 1


if __name__ == "__main__":
    sys.exit(main())
